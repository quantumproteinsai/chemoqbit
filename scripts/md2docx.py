#!/usr/bin/env python3
"""Render a manuscript Markdown file to .docx.

The sibling of scripts/md2pdf.py, for the journals and preprint servers that
want Word rather than PDF. It parses the same subset of Markdown, block for
block, so the two outputs say the same thing: YAML front matter, headings,
tables, fenced code, bullet and numbered lists, figures, and inline bold,
italic, code and links.

One difference is worth knowing. The PDF route has to borrow a second typeface
for Greek, arrows and angle brackets, because the built-in PDF fonts cannot draw
them. Word carries Unicode natively, so nothing is borrowed here and nothing can
come out as a black square. The super- and subscript characters are still
converted to real superscript and subscript runs rather than left as glyphs,
since that is what a copy-editor expects to find.

    python scripts/md2docx.py paper.md paper.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

# Unicode super- and subscripts become real Word runs. The two modifier capitals
# are here because the manuscripts use them where a subscript is meant -- the
# Boltzmann constant and the S_N2 label -- and a reader of the Word file should
# be able to search for "SN2".
SUP = dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺',
                '0123456789-+'))
SUB = dict(zip('₀₁₂₃₄₅₆₇₈₉₋'
               'ᴮᴺ',
                '0123456789-' 'BN'))
# Two characters that are plain mistakes in the sources rather than notation: a
# box-drawing rule standing in for an em dash in a journal title, and U+2010 for
# a hyphen. Word would render both, so this is tidying, not a workaround.
PLAIN = {'‐': '-', '─': '—'}

TEXT_CM = 16.6          # A4 less the 2.2 cm margins set below
BODY_PT = 10
FONT = 'Times New Roman'
MONO = 'Consolas'

TOKEN = re.compile(r'(\*\*[^*]+\*\*'
                   r'|(?<!\*)\*[^*]+\*(?!\*)'
                   r'|`[^`]+`'
                   r'|\[[^\]]+\]\([^)]+\))')


def _pieces(t):
    """Markdown inline -> [(text, flags)], flags drawn from b/i/code/small."""
    out = []
    for part in TOKEN.split(t):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            out.append((part[2:-2], {'b'}))
        elif part.startswith('`') and part.endswith('`'):
            out.append((part[1:-1], {'code'}))
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', part):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            out.append((m.group(1), set()))
            out.append((' <%s>' % m.group(2), {'small'}))
        elif part.startswith('*') and part.endswith('*'):
            out.append((part[1:-1], {'i'}))
        else:
            out.append((part, set()))
    return out


def _split_scripts(text):
    """Break a string into (text, 'sup'|'sub'|None) segments."""
    seg, buf, i, n = [], '', 0, len(text)
    while i < n:
        ch = text[i]
        tbl, kind = (SUP, 'sup') if ch in SUP else (SUB, 'sub') if ch in SUB else (None, None)
        if tbl is None:
            buf += PLAIN.get(ch, ch)
            i += 1
            continue
        if buf:
            seg.append((buf, None)); buf = ''
        run = ''
        while i < n and text[i] in tbl:
            run += tbl[text[i]]; i += 1
        seg.append((run, kind))
    if buf:
        seg.append((buf, None))
    return seg


def add_runs(p, text, size=BODY_PT, base=()):
    """Append the inline markup of `text` to paragraph `p`."""
    for chunk, flags in _pieces(text):
        flags = set(flags) | set(base)
        for s, script in _split_scripts(chunk):
            if not s:
                continue
            r = p.add_run(s)
            r.font.name = MONO if 'code' in flags else FONT
            r.font.size = Pt(size - 1.5 if 'code' in flags or 'small' in flags else size)
            r.bold = 'b' in flags
            r.italic = 'i' in flags
            if script == 'sup':
                r.font.superscript = True
            elif script == 'sub':
                r.font.subscript = True
            if 'small' in flags:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p


def _shade(cell, hex_fill):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def _rule(doc):
    """A horizontal rule as a paragraph bottom border, not a table."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:color'), 'B0B0B0')
    bottom.set(qn('w:space'), '1')
    borders.append(bottom)
    pPr.append(borders)
    return p


def _table(doc, rows):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    width = Cm(TEXT_CM / cols)
    for n_row, row in enumerate(rows):
        cells = t.add_row().cells
        for k in range(cols):
            cells[k].width = width
            p = cells[k].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, row[k] if k < len(row) else '', size=8.5,
                     base=('b',) if n_row == 0 else ())
            if n_row == 0:
                _shade(cells[k], 'EDEDED')
    return t


def _figure(doc, path, src):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f = Path(src).parent / path
    if not f.exists():
        add_runs(p, '*[missing figure: %s]*' % path)
        return p
    p.add_run().add_picture(str(f), width=Cm(TEXT_CM))
    return p


def convert(src, dst):
    lines = Path(src).read_text().split('\n')
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.2)

    i, n = 0, len(lines)

    # YAML front matter, as the JOSS paper carries
    if lines and lines[0].strip() == '---':
        j = next((k for k in range(1, n) if lines[k].strip() == '---'), None)
        if j:
            for ln in lines[1:j]:
                m = re.match(r'^(title|authors?|date|tags|bibliography):\s*(.*)$', ln)
                if m and m.group(2).strip():
                    v = m.group(2).strip().strip('"\'')
                    if m.group(1) == 'title':
                        add_runs(doc.add_heading('', level=1), v, size=16)
                    else:
                        add_runs(doc.add_paragraph(), v, size=9.5, base=('i',))
                elif re.match(r'^\s+-?\s*name:\s*(.+)$', ln):
                    add_runs(doc.add_paragraph(),
                             re.sub(r'^\s+-?\s*name:\s*', '', ln), size=9.5, base=('i',))
            i = j + 1

    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1; continue

        if s.startswith('```'):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(b)
                r.font.name = MONO
                r.font.size = Pt(8.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        if re.match(r'^(---|\*\*\*|___)\s*$', s):
            _rule(doc); i += 1; continue

        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', s)
        if m:
            _figure(doc, m.group(2), src); i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading('', level=min(lvl, 4))
            add_runs(h, m.group(2), size=16 if lvl == 1 else 12.5 if lvl == 2 else 11,
                     base=('i',) if lvl >= 3 else ())
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1; continue

        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:|-]+\|?\s*$',
                                                        lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                r = lines[i].strip().strip('|')
                if not re.match(r'^[\s:|-]+$', r):
                    rows.append([c.strip() for c in r.split('|')])
                i += 1
            _table(doc, rows)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        m = re.match(r'^[-*]\s+(.*)$', s)
        if m:
            add_runs(doc.add_paragraph(style='List Bullet'), m.group(1))
            i += 1; continue

        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            add_runs(doc.add_paragraph(style='List Number'), m.group(2))
            i += 1; continue

        buf = []
        while i < n and lines[i].strip() and not re.match(
                r'^(#{1,4}\s|[-*]\s|\d+\.\s|\||```|!\[|---\s*$)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, ' '.join(buf))

    doc.save(str(dst))
    return dst


if __name__ == '__main__':
    out = convert(sys.argv[1], sys.argv[2])
    print('  %s -> %s (%d kB)' % (sys.argv[1], out, Path(out).stat().st_size // 1024))
